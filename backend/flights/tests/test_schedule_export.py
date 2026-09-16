"""Выгрузка суточного плана `[ТЗ 3.1.1]`.

Проверяется одно: файл совпадает с тем, что видно на экране. Выгрузка,
игнорирующая отбор, приносит не ту таблицу, за которой её открывали, —
и заметно это становится, когда по ней уже приняли решение.
"""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from io import BytesIO
from typing import TYPE_CHECKING, Any

import pytest
from rest_framework import status

from accounts.models import Role
from catalog.models import AircraftCategory, AircraftType, Airport
from core.services import storage
from fleet.models import Aircraft
from flights.models import Flight
from flights.services import export as schedule_export

if TYPE_CHECKING:
    from collections.abc import Callable

    from rest_framework.test import APIClient

    from accounts.models import Organization
    from counterparties.models import Client

pytestmark = pytest.mark.django_db

URL = "/api/v1/flights/export"


@pytest.fixture
def airports(db: None) -> None:
    for icao, iata, name in (("UUWW", "VKO", "Внуково"), ("ULLI", "LED", "Пулково")):
        Airport.objects.create(
            icao=icao,
            iata=iata,
            name_ru=name,
            name_en=name,
            country="RU",
            timezone="Europe/Moscow",
            lat=55.0,
            lon=37.0,
        )


@pytest.fixture
def aircraft(db: None) -> Aircraft:
    aircraft_type = AircraftType.objects.create(
        icao_type="GLF5",
        name_ru="Gulfstream G550",
        name_en="Gulfstream G550",
        category=AircraftCategory.HEAVY,
        seats=16,
        cruise_speed_kts=488,
        fuel_burn_kg_per_hour=1300,
        turnaround_min=60,
    )
    return Aircraft.objects.create(
        registration="RA-10600", type=aircraft_type, home_base_icao="UUWW"
    )


@pytest.fixture
def flights(
    organization: Organization,
    client_alpha: Client,
    client_beta: Client,
    aircraft: Aircraft,
    airports: None,
) -> list[Flight]:
    departure = datetime(2026, 10, 5, 8, 0, tzinfo=UTC)
    return [
        Flight.objects.create(
            organization=organization,
            number=f"SLG-50{index:02d}",
            client=client_alpha if index % 2 == 0 else client_beta,
            aircraft=aircraft,
            type="charter",
            dep_icao="UUWW",
            arr_icao="ULLI",
            std_utc=departure + timedelta(days=index),
            sta_utc=departure + timedelta(days=index, hours=2),
            pax_count=4 + index,
            billing_currency="RUB",
        )
        for index in range(4)
    ]


@pytest.fixture(autouse=True)
def storage_stub(monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    stored: dict[str, Any] = {}

    def fake_put(*, key: str, data: bytes, mime_type: str) -> None:
        stored["key"] = key
        stored["data"] = data

    monkeypatch.setattr(storage, "put_bytes", fake_put)
    monkeypatch.setattr(
        storage, "presign_get", lambda *, key, file_name: f"http://storage.test/{key}"
    )
    return stored


def sheet_of(data: bytes) -> Any:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data))
    sheet = workbook.active
    assert sheet is not None
    return sheet


def test_export_returns_a_download_link(
    as_role: Callable[..., APIClient], flights: list[Flight], storage_stub: dict[str, Any]
) -> None:
    api = as_role(Role.DISPATCHER)
    response = api.post(URL, {}, format="json")

    assert response.status_code == status.HTTP_202_ACCEPTED, response.data
    body = response.json()
    assert body["status"] == "ready"
    assert body["downloadUrl"].startswith("http://storage.test/exports/")
    assert storage_stub["key"].endswith(".xlsx")


def test_file_holds_every_flight_of_the_selection(
    as_role: Callable[..., APIClient], flights: list[Flight], storage_stub: dict[str, Any]
) -> None:
    api = as_role(Role.DISPATCHER)
    api.post(URL, {}, format="json")

    values = [
        cell.value for row in sheet_of(storage_stub["data"]).iter_rows() for cell in row
    ]
    for flight in flights:
        assert flight.number in values


def test_filter_reaches_the_file(
    as_role: Callable[..., APIClient],
    flights: list[Flight],
    client_alpha: Client,
    storage_stub: dict[str, Any],
) -> None:
    """Выгрузка без отбора приносит не ту таблицу, за которой её открывали."""
    api = as_role(Role.DISPATCHER)
    api.post(f"{URL}?clientId={client_alpha.pk}", {}, format="json")

    values = [
        cell.value for row in sheet_of(storage_stub["data"]).iter_rows() for cell in row
    ]
    own = [flight for flight in flights if flight.client_id == client_alpha.pk]
    others = [flight for flight in flights if flight.client_id != client_alpha.pk]

    assert own, "приспособление не дало ни одного рейса этого клиента"
    for flight in own:
        assert flight.number in values
    for flight in others:
        assert flight.number not in values


def test_search_filter_reaches_the_file(
    as_role: Callable[..., APIClient], flights: list[Flight], storage_stub: dict[str, Any]
) -> None:
    api = as_role(Role.DISPATCHER)
    api.post(f"{URL}?search={flights[0].number}", {}, format="json")

    values = [
        cell.value for row in sheet_of(storage_stub["data"]).iter_rows() for cell in row
    ]
    assert flights[0].number in values
    assert flights[1].number not in values


def test_totals_are_the_sum_of_the_rows(
    as_role: Callable[..., APIClient], flights: list[Flight], storage_stub: dict[str, Any]
) -> None:
    """«Итого» обязано сходиться с колонкой (ADR-002 п. 3)."""
    api = as_role(Role.DISPATCHER)
    api.post(URL, {}, format="json")

    numbers = [
        cell.value
        for row in sheet_of(storage_stub["data"]).iter_rows()
        for cell in row
        if isinstance(cell.value, int | float)
    ]
    assert sum(flight.pax_count for flight in flights) in numbers


def test_demo_stand_marks_the_file(
    as_role: Callable[..., APIClient],
    flights: list[Flight],
    storage_stub: dict[str, Any],
    settings: Any,
) -> None:
    """`SPEC.md § 8.6`: выгрузка со стенда опознаётся как выгрузка со стенда."""
    settings.DEMO_DATA = True
    api = as_role(Role.DISPATCHER)
    api.post(URL, {}, format="json")

    first = sheet_of(storage_stub["data"]).cell(row=1, column=1).value
    assert first is not None and "DEMO" in str(first)


def test_columns_match_the_screen(flights: list[Flight]) -> None:
    """Колонки файла — те же, что в таблице суточного плана."""
    keys = [column.key for column in schedule_export.SCHEDULE.columns]
    assert keys[:4] == ["number", "client", "aircraft", "type"]
    assert "status" in keys


@pytest.mark.parametrize("role", [Role.CLIENT, Role.VENDOR])
def test_portal_roles_cannot_export_the_schedule(
    as_role: Callable[..., APIClient], role: str, client_alpha: Client, vendor_alpha: Any
) -> None:
    """Выгрузка всего расписания порталам не отдаётся вовсе."""
    extra: dict[str, Any] = (
        {"client": client_alpha} if role == Role.CLIENT else {"vendor": vendor_alpha}
    )
    api = as_role(role, **extra)
    assert api.post(URL, {}, format="json").status_code == status.HTTP_403_FORBIDDEN
