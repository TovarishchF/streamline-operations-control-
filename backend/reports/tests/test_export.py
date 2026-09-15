"""Выгрузка отчётов в четырёх форматах `[ТЗ 3.6.3]` (`SPEC.md § 9.3`).

Главное требование — то же, что у счёта: число в файле и число на экране
совпадают. Поэтому выгрузка строится из готового результата отчёта
и ничего не пересчитывает, а тесты сверяют файл с этим результатом,
а не с базой.

XML проверяется против `soc-report-v1.xsd` (ADR-032): схема объявлена
в решении, значит она обязана существовать и совпадать с выгрузкой.
"""

from __future__ import annotations

import csv
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from typing import TYPE_CHECKING, Any

import pytest
from django.conf import settings as django_settings
from rest_framework import status

from accounts.models import Role
from core.services import storage
from reports import definitions
from reports.services import builders, export

if TYPE_CHECKING:
    from collections.abc import Callable

    from rest_framework.test import APIClient

    from billing.models import Invoice
    from orders.models import ServiceOrder

pytestmark = pytest.mark.django_db

REPORTS_URL = "/api/v1/reports"
XSD_PATH = Path(django_settings.SHARED_DIR) / "reference" / "xsd" / "soc-report-v1.xsd"


def weasyprint_available() -> bool:
    """Доступна ли сборка PDF в этом окружении.

    WeasyPrint опирается на Pango и Cairo. В образе сервера они стоят
    (`backend/Dockerfile`), на машине разработчика под Windows — как
    правило нет. Пропустить проверку честнее, чем подменить её заглушкой
    и отчитаться, что PDF собирается.
    """
    try:
        import weasyprint  # noqa: F401
    except (ImportError, OSError):
        return False
    return True


requires_pdf = pytest.mark.skipif(
    not weasyprint_available(),
    reason="WeasyPrint требует Pango и Cairo; в образе сервера они есть, локально — нет",
)


def idempotent(key: str = "rep-key-00000001") -> dict[str, str]:
    return {"Idempotency-Key": key}


@pytest.fixture
def financial_result(
    completed_order: ServiceOrder,
    issued_invoice: Invoice,
    period: dict[str, Any],
    settings: Any,
) -> dict[str, Any]:
    """Отчёт боевого режима.

    Режим задаётся явно, а не наследуется: в демонстрационном файл
    начинается со строки с пометкой, и все проверки положения ячеек
    разъехались бы на строку. Разметку стенда проверяют отдельные
    тесты ниже.
    """
    settings.DEMO_DATA = False
    return builders.build(definitions.FINANCIAL, period)


@pytest.fixture
def spec() -> Any:
    return definitions.definition(definitions.FINANCIAL)


# ─────────────────────────── CSV ───────────────────────────


def test_csv_opens_in_excel_with_cyrillic(
    spec: Any, financial_result: dict[str, Any]
) -> None:
    """BOM и точка с запятой: без них русская локаль Excel портит файл."""
    data = export._render_csv(spec, financial_result)

    assert data.startswith(b"\xef\xbb\xbf")
    text = data.decode("utf-8-sig")
    header = text.splitlines()[0]
    assert header.split(";")[0] == "Рейс"


def test_csv_carries_the_same_numbers_as_the_report(
    spec: Any, financial_result: dict[str, Any]
) -> None:
    text = export._render_csv(spec, financial_result).decode("utf-8-sig")
    rows = list(csv.reader(StringIO(text), delimiter=";"))

    keys = [column.key for column in spec.columns]
    revenue_at = keys.index("revenue")
    data_row = rows[1]

    assert Decimal(data_row[revenue_at]) == Decimal(financial_result["rows"][0]["revenue"])

    totals_row = rows[-1]
    assert totals_row[0] == "Итого"
    assert Decimal(totals_row[revenue_at]) == Decimal(financial_result["totals"]["revenue"])


def test_csv_leaves_missing_value_empty_not_zero(spec: Any) -> None:
    """Пустая ячейка честнее нуля: величина не посчитана, а не равна нулю."""
    result = {
        "code": definitions.FINANCIAL,
        "generatedAt": "2026-09-15T00:00:00+00:00",
        "currency": "RUB",
        "isDemo": False,
        "columns": spec.columns_to_contract(),
        "rows": [
            {
                "flight": "SLG-1",
                "client": "Клиент",
                "revenue": "0.00",
                "cost": "0.00",
                "margin": "0.00",
                "marginPercent": None,
            }
        ],
        "totals": {},
    }
    text = export._render_csv(spec, result).decode("utf-8-sig")
    rows = list(csv.reader(StringIO(text), delimiter=";"))
    assert rows[1][-1] == ""


# ─────────────────────────── XLSX ───────────────────────────


def test_xlsx_money_cells_are_numbers(spec: Any, financial_result: dict[str, Any]) -> None:
    """Суммы в таблице должны складываться средствами Excel."""
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(export._render_xlsx(spec, financial_result)))
    sheet = workbook.active
    assert sheet is not None

    keys = [column.key for column in spec.columns]
    revenue_column = keys.index("revenue") + 1
    cell = sheet.cell(row=2, column=revenue_column)

    assert isinstance(cell.value, int | float)
    assert Decimal(str(cell.value)).quantize(Decimal("0.01")) == Decimal(
        financial_result["rows"][0]["revenue"]
    )


def test_xlsx_totals_match_the_report(spec: Any, financial_result: dict[str, Any]) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(export._render_xlsx(spec, financial_result)))
    sheet = workbook.active
    assert sheet is not None

    numbers = [
        Decimal(str(cell.value)).quantize(Decimal("0.01"))
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, int | float)
    ]
    expected = Decimal(financial_result["totals"]["revenue"])
    assert expected in numbers, f"итога {expected} нет в выгрузке: {numbers}"


# ─────────────────────────── XML ───────────────────────────


def test_xml_validates_against_soc_report_xsd(
    spec: Any, financial_result: dict[str, Any], period: dict[str, Any]
) -> None:
    """ADR-032: выгрузка проверяется против объявленной схемы."""
    from lxml import etree

    assert XSD_PATH.exists(), f"схема {XSD_PATH} объявлена в ADR-032, но отсутствует"
    schema = etree.XMLSchema(etree.parse(str(XSD_PATH)))

    document = etree.fromstring(export._render_xml(spec, financial_result, period))
    schema.assertValid(document)


@pytest.mark.parametrize("code", sorted(definitions.DEFINITIONS))
def test_every_report_validates_against_the_schema(code: str, period: dict[str, Any]) -> None:
    """Схема одна на все отчёты — значит, каждый обязан ей соответствовать."""
    from lxml import etree

    schema = etree.XMLSchema(etree.parse(str(XSD_PATH)))
    parameters = period if code != definitions.RECEIVABLES_PAYABLES else {}
    result = builders.build(code, parameters)

    document = etree.fromstring(
        export._render_xml(definitions.definition(code), result, parameters)
    )
    schema.assertValid(document)


def test_xml_money_carries_its_currency(
    spec: Any, financial_result: dict[str, Any], period: dict[str, Any]
) -> None:
    """Сумма без валюты в выгрузке бессмысленна: отчёт бывает многовалютным."""
    from lxml import etree

    document = etree.fromstring(export._render_xml(spec, financial_result, period))
    namespace = {"r": export.XML_NAMESPACE}

    revenue = document.find(".//r:Rows/r:Row/r:Value[@key='revenue']", namespace)
    assert revenue is not None
    assert revenue.get("currency") == "RUB"
    assert revenue.text == financial_result["rows"][0]["revenue"]


def test_xml_records_the_parameters_it_was_built_with(
    spec: Any, financial_result: dict[str, Any], period: dict[str, Any]
) -> None:
    from lxml import etree

    document = etree.fromstring(export._render_xml(spec, financial_result, period))
    namespace = {"r": export.XML_NAMESPACE}

    values = {
        element.get("key"): element.text
        for element in document.findall(".//r:Parameters/r:Parameter", namespace)
    }
    assert values["from"] == period["from"].isoformat()
    assert values["to"] == period["to"].isoformat()


# ─────────────────────────── PDF ───────────────────────────


@requires_pdf
def test_pdf_is_produced(spec: Any, financial_result: dict[str, Any]) -> None:
    data = export._render_pdf(spec, financial_result)
    assert data.startswith(b"%PDF-")
    assert len(data) > 1000


# ─────────────────────────── Разметка стенда ───────────────────────────


def test_demo_stand_marks_every_format(
    spec: Any, financial_result: dict[str, Any], period: dict[str, Any]
) -> None:
    """ADR-008: стенд опознаётся и там, где водяной знак нарисовать негде."""
    from lxml import etree
    from openpyxl import load_workbook

    financial_result["isDemo"] = True

    assert export.DEMO_NOTE.encode() in export._render_csv(spec, financial_result)

    workbook = load_workbook(BytesIO(export._render_xlsx(spec, financial_result)))
    sheet = workbook.active
    assert sheet is not None
    assert "DEMO" in str(sheet.cell(row=1, column=1).value)

    document = etree.fromstring(export._render_xml(spec, financial_result, period))
    assert document.get("demo") == "true"
    note = document.find(f"{{{export.XML_NAMESPACE}}}Note")
    assert note is not None and note.text == export.DEMO_NOTE


def test_production_export_has_no_demo_marking(
    spec: Any, financial_result: dict[str, Any], period: dict[str, Any]
) -> None:
    """В боевом режиме маркировки нет (`CLAUDE.md § 4`)."""
    from lxml import etree

    financial_result["isDemo"] = False

    assert export.DEMO_NOTE.encode() not in export._render_csv(spec, financial_result)
    document = etree.fromstring(export._render_xml(spec, financial_result, period))
    assert document.get("demo") == "false"
    assert document.find(f"{{{export.XML_NAMESPACE}}}Note") is None


# ─────────────────────────── Эндпоинт ───────────────────────────


@pytest.fixture
def storage_stub(monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    """Хранилище подменяется: проверяется выгрузка, а не MinIO."""
    stored: dict[str, Any] = {}

    def fake_put(*, key: str, data: bytes, mime_type: str) -> None:
        stored["key"] = key
        stored["size"] = len(data)
        stored["mime"] = mime_type

    monkeypatch.setattr(storage, "put_bytes", fake_put)
    monkeypatch.setattr(
        storage, "presign_get", lambda *, key, file_name: f"http://storage.test/{key}"
    )
    return stored


@pytest.mark.parametrize("export_format", ["csv", "xlsx", "xml"])
def test_export_endpoint_returns_download_link(
    as_role: Callable[..., APIClient],
    storage_stub: dict[str, Any],
    completed_order: ServiceOrder,
    query: dict[str, str],
    export_format: str,
) -> None:
    api = as_role(Role.FINANCE)
    response = api.post(
        f"{REPORTS_URL}/{definitions.FINANCIAL}/export?from={query['from']}&to={query['to']}",
        {"format": export_format},
        format="json",
        headers=idempotent(f"rep-{export_format}-0001"),
    )

    assert response.status_code == status.HTTP_202_ACCEPTED, response.data
    body = response.json()
    assert body["status"] == "ready"
    assert body["downloadUrl"].startswith("http://storage.test/reports/")
    assert storage_stub["key"].endswith(f".{export_format}")
    assert storage_stub["size"] > 0


def test_export_of_unknown_report_is_not_found(
    as_role: Callable[..., APIClient], storage_stub: dict[str, Any]
) -> None:
    api = as_role(Role.FINANCE)
    response = api.post(
        f"{REPORTS_URL}/margin_by_moon_phase/export",
        {"format": "csv"},
        format="json",
        headers=idempotent("rep-none-00001"),
    )
    assert response.status_code == status.HTTP_404_NOT_FOUND


def test_unsupported_format_is_refused(
    as_role: Callable[..., APIClient], storage_stub: dict[str, Any]
) -> None:
    api = as_role(Role.FINANCE)
    response = api.post(
        f"{REPORTS_URL}/{definitions.FINANCIAL}/export",
        {"format": "docx"},
        format="json",
        headers=idempotent("rep-bad-000001"),
    )
    assert response.status_code == status.HTTP_400_BAD_REQUEST


def test_export_requires_idempotency_key(
    as_role: Callable[..., APIClient], storage_stub: dict[str, Any]
) -> None:
    """ADR-018: повторное нажатие кнопки не должно плодить файлы."""
    api = as_role(Role.FINANCE)
    response = api.post(
        f"{REPORTS_URL}/{definitions.FINANCIAL}/export", {"format": "csv"}, format="json"
    )
    assert response.status_code == status.HTTP_400_BAD_REQUEST


def test_dispatcher_cannot_export_financial_report(
    as_role: Callable[..., APIClient], storage_stub: dict[str, Any]
) -> None:
    """Запрет на просмотр обязан распространяться на выгрузку."""
    api = as_role(Role.DISPATCHER)
    response = api.post(
        f"{REPORTS_URL}/{definitions.FINANCIAL}/export",
        {"format": "csv"},
        format="json",
        headers=idempotent("rep-disp-00001"),
    )
    assert response.status_code == status.HTTP_403_FORBIDDEN
