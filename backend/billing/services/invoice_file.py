"""Файл счёта поставщика для сверки `[ТЗ 3.4.2]`.

Шаблон и разбор описаны здесь вместе, одним перечнем колонок: шаблон,
который система выдаёт, и файл, который она умеет читать, разойтись
не могут.

Файл — одна таблица, строка на позицию счёта. Шапка счёта (номер, дата,
валюта) заполняется в первой строке; в остальных её можно оставить пустой,
а заполненная обязана совпадать с первой: файл с двумя счетами — ошибка,
а не два импорта.

Разбор не сверяет. Он приводит файл к той же форме `VendorInvoiceImport`,
что приходит JSON-ом, и дальше импорт идёт одним путём. Ошибки собираются
по всему файлу с номерами строк: исправлять файл по одной ошибке за попытку —
худший способ узнать, что с ним не так.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePath
from typing import TYPE_CHECKING, Any

from django.conf import settings
from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException

from core.exceptions import DomainError

if TYPE_CHECKING:
    from collections.abc import Callable

    Cell = Callable[[str], Any]

# Счёт на несколько тысяч позиций занимает сотни килобайт; файл больше
# этого — не счёт.
MAX_BYTES = 5 * 1024 * 1024

# Больше ошибок перечислять бессмысленно: файл составлен не по шаблону.
MAX_ERRORS = 20

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_CONTENT_TYPE = "text/csv; charset=utf-8"

ICAO = re.compile(r"^[A-Za-z]{4}$")

DATE_FORMATS = (
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y %H:%M",
    "%d.%m.%Y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
)


class _Semicolon(csv.excel):
    """CSV из Excel с русской локалью: разделитель — точка с запятой."""

    delimiter = ";"


class InvoiceFileError(DomainError):
    """Файл счёта не разобран. Подробности — в `details.rows`."""

    code = "VALIDATION_ERROR"


@dataclass(frozen=True, slots=True)
class Column:
    """Колонка шаблона: поле `VendorInvoiceImport` и как его заполнять."""

    key: str
    title: str
    hint: str
    example: str


HEADER_COLUMNS = (
    Column("number", "Номер счёта", "Текст. Заполняется в первой строке", "А-1024"),
    Column("issuedAt", "Дата счёта", "ДД.ММ.ГГГГ. Заполняется в первой строке", "20.10.2026"),
    Column("currency", "Валюта", "RUB, USD или EUR. Заполняется в первой строке", "RUB"),
)

LINE_COLUMNS = (
    Column("airportIcao", "Аэропорт ИКАО", "Четыре латинские буквы", "UUWW"),
    Column(
        "serviceDate",
        "Дата услуги",
        "ДД.ММ.ГГГГ, можно с временем ЧЧ:ММ. Время — UTC",
        "18.10.2026 14:30",
    ),
    Column(
        "serviceCode",
        "Код услуги",
        "Код из номенклатуры поставщика или код из каталога системы",
        "FUEL-JETA1",
    ),
    Column("quantity", "Количество", "Число, дробная часть через запятую", "1000"),
    Column("unitPrice", "Цена за единицу", "Число, дробная часть через запятую", "70,00"),
    Column("amount", "Сумма", "Число, дробная часть через запятую", "70000,00"),
)

COLUMNS = HEADER_COLUMNS + LINE_COLUMNS


def _normalized(text: str) -> str:
    return " ".join(text.replace("ё", "е").replace("Ё", "Е").lower().split())


# Заголовок узнаётся по названию из шаблона и по имени поля API: второе
# удобно, когда файл выгружает система поставщика, а не человек.
ALIASES = {
    alias: column.key
    for column in COLUMNS
    for alias in (_normalized(column.title), _normalized(column.key))
}


# ─────────────────────────── Шаблон ───────────────────────────


def template(file_format: str) -> tuple[bytes, str, str]:
    """Шаблон файла: содержимое, тип и имя для сохранения."""
    if file_format == "csv":
        buffer = io.StringIO()
        # Точка с запятой и BOM — так файл откроется в Excel с русской
        # локалью по колонкам, а не одной строкой.
        csv.writer(buffer, delimiter=";").writerow([column.title for column in COLUMNS])
        return (
            buffer.getvalue().encode("utf-8-sig"),
            CSV_CONTENT_TYPE,
            "soc-vendor-invoice-template.csv",
        )

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None  # у новой книги лист есть всегда
    sheet.title = _("Счёт")
    sheet.append([column.title for column in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for letter, column in zip("ABCDEFGHI", COLUMNS, strict=True):
        sheet.column_dimensions[letter].width = max(16, len(column.title) + 4)
    sheet.freeze_panes = "A2"

    guide = workbook.create_sheet(_("Инструкция"))
    bold_rows: list[int] = []

    def add(values: list[str], *, bold: bool = False) -> None:
        guide.append(values)
        if bold:
            bold_rows.append(guide.max_row)

    add([_("Как заполнить")], bold=True)
    add([_("Строка на позицию счёта. Читается только первый лист, этот лист — справка.")])
    add([_("Номер, дату и валюту счёта достаточно указать в первой строке.")])
    add([])
    add([_("Колонка"), _("Формат"), _("Пример")], bold=True)
    for column in COLUMNS:
        add([column.title, column.hint, column.example])
    add([])
    add([_("Пример заполнения первого листа")], bold=True)
    add([column.title for column in COLUMNS], bold=True)
    add([column.example for column in COLUMNS])
    add([""] * len(HEADER_COLUMNS) + ["UUWW", "18.10.2026 15:10", "GPU", "2", "4500,00", "9000,00"])
    for row in bold_rows:
        for cell in guide[row]:
            cell.font = Font(bold=True)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 60
    guide.column_dimensions["C"].width = 20

    buffer_bytes = io.BytesIO()
    workbook.save(buffer_bytes)
    return buffer_bytes.getvalue(), XLSX_CONTENT_TYPE, "soc-vendor-invoice-template.xlsx"


# ─────────────────────────── Разбор ───────────────────────────


def parse(name: str, content: bytes) -> dict[str, Any]:
    """Приводит файл счёта к форме `VendorInvoiceImport` без `vendorId`."""
    if len(content) > MAX_BYTES:
        raise InvoiceFileError(
            _("Файл больше %(size)s МБ — это не счёт") % {"size": MAX_BYTES // 1024 // 1024},
            {"rows": []},
        )

    suffix = PurePath(name).suffix.lower()
    if suffix == ".xlsx":
        rows = _xlsx_rows(content)
    elif suffix in (".csv", ".txt"):
        rows = _csv_rows(content)
    else:
        raise InvoiceFileError(
            _("Формат файла не поддерживается: нужен XLSX или CSV по шаблону"), {"rows": []}
        )
    return _to_payload(rows)


def _xlsx_rows(content: bytes) -> list[list[Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as error:
        raise InvoiceFileError(
            _("Файл XLSX не читается: возможно, он повреждён или сохранён в другом формате"),
            {"rows": []},
        ) from error
    try:
        # Читается только первый лист: в шаблоне второй — инструкция.
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _csv_rows(content: bytes) -> list[list[Any]]:
    # Excel с русской локалью сохраняет CSV в Windows-1251, остальные — в UTF-8.
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1251", errors="replace")

    dialect: type[csv.Dialect]
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = _Semicolon
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _to_payload(rows: list[list[Any]]) -> dict[str, Any]:
    errors: list[str] = []

    # Заголовок — первая непустая строка.
    start = next((index for index, row in enumerate(rows) if not all(map(_blank, row))), None)
    if start is None:
        raise InvoiceFileError(_("Файл пустой"), {"rows": []})

    positions: dict[str, int] = {}
    for index, title in enumerate(rows[start]):
        key = ALIASES.get(_normalized(str(title))) if not _blank(title) else None
        if key is not None and key not in positions:
            positions[key] = index

    missing = [column.title for column in COLUMNS if column.key not in positions]
    if missing:
        raise InvoiceFileError(
            _("В файле нет колонок: %(columns)s. Скачайте шаблон и заполните его")
            % {"columns": ", ".join(missing)},
            {"rows": [], "missingColumns": missing},
        )

    header: dict[str, str] = {}
    lines: list[dict[str, Any]] = []
    first = True

    for index in range(start + 1, len(rows)):
        row = rows[index]
        if all(map(_blank, row)):
            continue
        row_number = index + 1

        def cell(key: str, row: list[Any] = row) -> Any:
            position = positions[key]
            return row[position] if position < len(row) else None

        row_errors: list[str] = []
        _read_header(cell, header, row_errors, first=first)
        first = False
        line = _read_line(cell, row_errors)
        if row_errors:
            errors.extend(
                _("Строка %(row)s: %(error)s") % {"row": row_number, "error": error}
                for error in row_errors
            )
        elif line is not None:
            lines.append(line)

        if len(errors) >= MAX_ERRORS:
            break

    if not errors and not lines:
        errors.append(_("В файле нет ни одной позиции счёта"))
    if errors:
        raise InvoiceFileError(
            _("Файл не разобран: ошибок %(count)s") % {"count": len(errors)},
            {"rows": errors[:MAX_ERRORS]},
        )

    for line in lines:
        line["unitPrice"]["currency"] = header["currency"]
        line["amount"]["currency"] = header["currency"]
    return {**header, "lines": lines}


def _read_header(cell: Cell, header: dict[str, str], errors: list[str], *, first: bool) -> None:
    """Шапка счёта: из первой строки, в остальных — пусто или то же самое."""
    readers: dict[str, Callable[[Any], str]] = {
        "number": lambda value: str(value).strip(),
        "issuedAt": lambda value: _moment(value).isoformat(),
        "currency": _currency,
    }
    titles = {column.key: column.title for column in HEADER_COLUMNS}

    for key, read in readers.items():
        raw = cell(key)
        if _blank(raw):
            if first:
                errors.append(_("не заполнено поле «%(title)s»") % {"title": titles[key]})
            continue
        try:
            value = read(raw)
        except ValueError as error:
            errors.append(f"{titles[key]}: {error}")
            continue

        if first or key not in header:
            header[key] = value
        elif header[key] != value:
            errors.append(
                _(
                    "«%(title)s» %(value)s не совпадает с первой строкой (%(first)s): "
                    "в файле должен быть один счёт"
                )
                % {"title": titles[key], "value": value, "first": header[key]}
            )


def _read_line(cell: Cell, errors: list[str]) -> dict[str, Any] | None:
    """Позиция счёта в форме `VendorInvoiceLine`; валюта подставляется из шапки."""
    titles = {column.key: column.title for column in LINE_COLUMNS}
    values: dict[str, Any] = {}

    readers: dict[str, Callable[[Any], Any]] = {
        "airportIcao": _icao,
        "serviceDate": lambda value: _moment(value).isoformat(),
        "serviceCode": lambda value: str(value).strip(),
        "quantity": _decimal,
        "unitPrice": _decimal,
        "amount": _decimal,
    }
    for key, read in readers.items():
        raw = cell(key)
        if _blank(raw):
            errors.append(_("не заполнено поле «%(title)s»") % {"title": titles[key]})
            continue
        try:
            values[key] = read(raw)
        except ValueError as error:
            errors.append(f"{titles[key]}: {error}")

    if len(values) < len(readers):
        return None
    return {
        "airportIcao": values["airportIcao"],
        "serviceDate": values["serviceDate"],
        "serviceCode": values["serviceCode"],
        "quantity": str(values["quantity"]),
        "unitPrice": {"amount": str(values["unitPrice"]), "currency": ""},
        "amount": {"amount": str(values["amount"]), "currency": ""},
    }


def _icao(value: Any) -> str:
    text = str(value).strip().upper()
    if not ICAO.match(text):
        raise ValueError(_("«%(value)s» — не код ИКАО из четырёх букв") % {"value": value})
    return text


def _currency(value: Any) -> str:
    text = str(value).strip().upper()
    if text not in settings.SUPPORTED_CURRENCIES:
        raise ValueError(
            _("«%(value)s» — нужна одна из %(allowed)s")
            % {"value": value, "allowed": ", ".join(settings.SUPPORTED_CURRENCIES)}
        )
    return text


def _decimal(value: Any) -> Decimal:
    """Число из ячейки — сразу в `Decimal` (`CLAUDE.md § 3` п. 1)."""
    if isinstance(value, bool):
        raise ValueError(_("«%(value)s» — не число") % {"value": value})
    if isinstance(value, Decimal | int):
        return Decimal(value)
    if isinstance(value, float):
        # XLSX хранит числа двоичной плавающей точкой. `str()` даёт кратчайшую
        # запись, которая читается обратно в то же число, — ту, что человек
        # видит в ячейке. Арифметики над float нет: значение сразу Decimal.
        return Decimal(str(value))

    text = str(value).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        number = Decimal(text)
    except InvalidOperation as error:
        raise ValueError(_("«%(value)s» — не число") % {"value": value}) from error
    if not number.is_finite():
        raise ValueError(_("«%(value)s» — не число") % {"value": value})
    return number


def _moment(value: Any) -> datetime:
    """Дата или дата со временем; без пояса — UTC (`CLAUDE.md § 3` п. 2)."""
    if isinstance(value, datetime):
        return value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=UTC)

    text = str(value).strip()
    for pattern in DATE_FORMATS:
        try:
            return datetime.strptime(text, pattern).replace(tzinfo=UTC)
        except ValueError:
            continue
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError as error:
        raise ValueError(
            _("«%(value)s» — нужна дата ДД.ММ.ГГГГ или ГГГГ-ММ-ДД") % {"value": value}
        ) from error
    return parsed.replace(tzinfo=UTC) if parsed.tzinfo is None else parsed.astimezone(UTC)
